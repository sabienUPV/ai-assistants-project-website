import os
import glob
import re
import sys
import shutil
from urllib.request import urlopen
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

# --- Configuration ---
INPUT_DIR = "input"
LOCAL_IMG_DIR = os.path.join(INPUT_DIR, "extracted_images")

# Astro project relative paths
ASSETS_DIR = "../packages/admin/src/assets/images/courses"
CONTENT_DIR = "../packages/admin/src/content/en/courses"

LOGO_Y_THRESHOLD_TITLE = 0.15  # Upper 15% of the slide to filter out header logos in title slide (first slide)
LOGO_Y_THRESHOLD_CONTENT = 0.05  # Upper 5% of the slide to filter out header logos in content slides (subsequent slides)

# Configurable keywords to identify and skip redundant slides (case-insensitive)
SKIP_SLIDE_KEYWORDS = [
    "BY-NC-SA",
    "EACEA",
    "Consortium"
]

def extract_rich_text(shape):
    """
    Extracts text from a shape while preserving line breaks and 
    translating native PowerPoint bold formatting to Markdown.
    Merges consecutive bold runs to prevent fragmentation (e.g. **w** **o** **rd**).
    """
    if not getattr(shape, "has_text_frame", False) or not shape.has_text_frame:
        return ""
        
    paragraphs_text = []
    for p in shape.text_frame.paragraphs:
        p_text = ""
        merged_runs = []
        
        # 1. Merge consecutive runs with the same bold formatting
        for run in p.runs:
            if not run.text:
                continue
            
            is_bold = bool(run.font.bold)
            
            if merged_runs and merged_runs[-1]['bold'] == is_bold:
                merged_runs[-1]['text'] += run.text
            else:
                merged_runs.append({'text': run.text, 'bold': is_bold})
                
        # 2. Apply markdown formatting to the merged runs
        for item in merged_runs:
            text = item['text']
            if item['bold'] and text.strip():
                # Extract surrounding whitespace to place the markdown tags tightly around the word
                stripped = text.strip()
                left_space = text[:len(text)-len(text.lstrip())]
                right_space = text[len(text.rstrip()):]
                p_text += f"{left_space}**{stripped}**{right_space}"
            else:
                p_text += text
                
        cleaned_p_text = p_text.strip()
        if cleaned_p_text:
            # If the paragraph has a level > 0, it was likely a sub-list in PPT.
            # We add a hyphen to help the LLM identify it as a bullet point.
            if p.level > 0:
                cleaned_p_text = f"- {cleaned_p_text}"
            paragraphs_text.append(cleaned_p_text)
            
    return " \\n ".join(paragraphs_text)


def extract_pptx_data():
    """
    Parses PowerPoint files to extract text and physical images, 
    mapping their relative center (X, Y) spatial coordinates for LLM processing.
    Now acts as an interactive CLI tool for hybrid manual/automated generation.
    """
    # Find all .pptx files, ignoring Microsoft Office temporary lock files (~$)
    pptx_files = [
        f for f in glob.glob(os.path.join(INPUT_DIR, "*.pptx")) 
        if not os.path.basename(f).startswith("~$")
    ]

    if not pptx_files:
        print(f"⚠️ No valid .pptx files found in '{INPUT_DIR}'.", file=sys.stderr)
        return

    # Regex to enforce patterns like: "8 - Name.pptx", "8.2 - Name.pptx", "8.2.1 - Name.pptx", etc.
    filename_pattern = re.compile(r'^(?P<unit>\d+(?:\.\d+)*)\s*-\s*(?P<title>.+)\.pptx$')

    for pptx_path in pptx_files:
        filename = os.path.basename(pptx_path)
        
        match = filename_pattern.match(filename)
        if not match:
            print(f"⚠️ SKIPPED: '{filename}' does not match the 'Unit - Title.pptx' format (e.g. '8.2 - Social participation.pptx')", file=sys.stderr)
            continue
            
        unit_str = match.group('unit')      # e.g., "8.2"
        title_str = match.group('title')    # e.g., "Social participation and AI"
        unit_folder = unit_str.replace('.', '-') # e.g., "8-2"
        
        print(f"\n{'='*60}", file=sys.stderr)
        print(f"🚀 Processing: Unit {unit_str} - {title_str}", file=sys.stderr)
        print(f"{'='*60}", file=sys.stderr)
        
        # Local temp directory for images before post-processing
        output_img_dir = os.path.join(LOCAL_IMG_DIR, unit_folder)
        os.makedirs(output_img_dir, exist_ok=True)
        
        spatial_data_output = []

        try:
            prs = Presentation(pptx_path)
            slide_width = prs.slide_width
            slide_height = prs.slide_height

            for i, slide in enumerate(prs.slides):
                # Determine logo filter threshold based on slide index
                logo_threshold = LOGO_Y_THRESHOLD_TITLE if i == 0 else LOGO_Y_THRESHOLD_CONTENT
                
                # --- 1. PRE-SCAN: Check for skip keywords ---
                def get_full_slide_text(shape):
                    """Recursively aggregates all text from a shape and its grouped children."""
                    text = ""
                    if getattr(shape, 'shape_type', None) == MSO_SHAPE_TYPE.GROUP:
                        for sub_shape in shape.shapes:
                            text += get_full_slide_text(sub_shape)
                        return text
                    
                    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                        text += " ".join([p.text.strip() for p in shape.text_frame.paragraphs]) + " "
                    return text

                full_slide_text = ""
                for shape in slide.shapes:
                    full_slide_text += get_full_slide_text(shape)
                
                skip_slide_checks = [(kw, kw.lower() in full_slide_text.lower()) for kw in SKIP_SLIDE_KEYWORDS]
                
                if any([x[1] for x in skip_slide_checks]):
                    matched_keywords = [x[0] for x in skip_slide_checks if x[1]]
                    print(f"⏭️ Skipping slide {i+1} (Matched keywords: {matched_keywords})", file=sys.stderr)
                    continue
                # --------------------------------------------

                print(f"⚙️ Parsing Slide {i+1}...", file=sys.stderr)
                spatial_data_output.append(f"\n--- SLIDE {i + 1} ---")
                elements = []
                
                # --- 2. EXTRACTION: Process shapes and physical images ---
                def process_shape(shape):
                    """
                    Recursively processes shapes, extracting rich text, native pictures, 
                    and hidden blipFill image blobs.
                    """
                    # Handle grouped shapes by entering them recursively
                    if getattr(shape, 'shape_type', None) == MSO_SHAPE_TYPE.GROUP:
                        for sub_shape in shape.shapes:
                            process_shape(sub_shape)
                        return

                    # Ignore elements missing dimensional data
                    if not all(hasattr(shape, attr) for attr in ['top', 'left', 'width', 'height']):
                        return
                        
                    # Calculate center relative coordinates (0.0 to 1.0) for accurate alignment
                    center_x = (shape.left + (shape.width / 2)) / slide_width
                    center_y = (shape.top + (shape.height / 2)) / slide_height
                    top_y = shape.top / slide_height  # Kept strictly for the top-logo filter
                    
                    shape_type = getattr(shape, 'shape_type', None)
                    name = getattr(shape, 'name', 'Element').replace(" ", "_")
                    
                    text_content = extract_rich_text(shape)
                    
                    # LOGO FILTER: Ignore visual elements in the upper threshold (using absolute top)
                    if top_y < logo_threshold and not text_content:
                        return
                        
                    if text_content:
                        elements.append({
                            "type": "Text", 
                            "content": text_content, 
                            "x": center_x, 
                            "y": center_y
                        })
                    else:
                        blob = None
                        ext = "png"
                        
                        # 2A. Standard rasterized image (<p:pic>)
                        if shape_type == MSO_SHAPE_TYPE.PICTURE or hasattr(shape, "image"):
                            blob = shape.image.blob
                            ext = getattr(shape.image, "ext", "png")
                            
                        # 2B. Vector shape with a hidden image fill (Picture Fill) OR horizontal separator line
                        elif shape_type in [MSO_SHAPE_TYPE.FREEFORM, MSO_SHAPE_TYPE.AUTO_SHAPE, MSO_SHAPE_TYPE.LINE]:
                            if hasattr(shape, "fill") and getattr(shape.fill, "type", None) == 6:
                                try:
                                    # Navigate XML tree to find the relation ID (rId)
                                    blips = shape._element.xpath('.//a:blip')
                                    if blips:
                                        rId = blips[0].get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
                                        if rId:
                                            # Extract raw physical file via SlidePart relationships
                                            rel = shape.part.rels[rId]
                                            image_part = rel.target_part
                                            blob = image_part.blob
                                            
                                            # Attempt to extract actual extension from content type
                                            c_type = getattr(image_part, "content_type", "")
                                            ext = c_type.split('/')[-1] if '/' in c_type else "png"
                                except Exception as e:
                                    print(f"⚠️ Could not extract physical image from {name}: {e}", file=sys.stderr)
                            else:
                                # Check if it's a horizontal line
                                rel_width = shape.width / slide_width
                                rel_height = shape.height / slide_height

                                # A horizontal line is usually wide (> 15% of slide) and thin (< 2% height)
                                if rel_width > 0.15 and rel_height < 0.02:
                                    elements.append({
                                        "type": "Separator", 
                                        "content": "---", 
                                        "x": center_x, 
                                        "y": center_y
                                    })
                                else:
                                    # Pure vector path without image fill, and not a horizontal line. We just include a placeholder with the type of element for the LLM to handle.
                                    elements.append({"type": "Vector", "content": f"[{name}]", "x": center_x, "y": center_y})
                        
                        # If a binary image was successfully extracted, save it locally first
                        if blob:
                            image_name = f"slide_{i+1}_{name}.{ext}"
                            image_path = os.path.join(output_img_dir, image_name)
                            
                            with open(image_path, "wb") as f:
                                f.write(blob)
                                
                            # Prepare clean Markdown image syntax for LLM context (paths resolved in post-processing)
                            elements.append({
                                "type": "Image", 
                                "content": f"![{image_name}]()", 
                                "x": center_x, 
                                "y": center_y
                            })

                for shape in slide.shapes:
                    process_shape(shape)

                # Sort spatially by center point: top-to-bottom (Y), then left-to-right (X)
                elements.sort(key=lambda e: (round(e["y"], 2), e["x"]))

                for el in elements:
                    spatial_data_output.append(f"[X:{el['x']:.2f}, Y:{el['y']:.2f}] {el['type']}: {el['content']}")

        except Exception as e:
            print(f"❌ Error processing {filename}: {e}", file=sys.stderr)
            continue
            
        # --- GENERATE PROMPT (stdout) ---
        schema_url = "https://ai4pid.eu/llms.txt"
        schema_fallback = "[MANUALLY PASTE THE CONTENT OF YOUR llms.txt HERE]"
        schema_found = False
        try:
            with urlopen(schema_url, timeout=10) as response:
                schema_content = response.read().decode("utf-8").strip()
            if not schema_content:
                raise ValueError("The llms.txt response was empty")
            schema_found = True
        except Exception as e:
            print(f"⚠️ Could not retrieve llms.txt ({e}); using the manual placeholder.", file=sys.stderr)
            schema_content = schema_fallback

        prompt_text = f"""System Role:
You are an expert frontend developer and content architect. Your task is to convert a spatial coordinate map extracted from PowerPoint slides into a structured Markdoc (.mdoc) file for an e-training platform.

Context & Schema:
Below is the schema and custom tags you must use. 

<schema>
{schema_content}
</schema>

Spatial Mapping Rules:
You will receive raw data representing elements on slides. Each element has an [X, Y] coordinate representing its geometric center (0.0 to 1.0). 
* X=0.0 is the far left, X=1.0 is the far right. 
* Y=0.0 is the top, Y=1.0 is the bottom.

You must apply these logical rules to build the Markdoc layout:

1. Slideshow Wrapping: Wrap the entire output in a single {{% slideshow %}} component. Each `--- SLIDE X ---` must be wrapped in a {{% slide title="[Extracted Title]" %}} tag. The title is usually the text element with the lowest Y value.
2. Column Layouts (The X Axis): If you detect elements positioned on opposite sides of the screen (e.g., an image at X:0.10 and text at X:0.50), you MUST use the {{% columns %}} and {{% column %}} tags to recreate the visual layout.
3. Paragraph Splitting (The Y Axis) - CRITICAL RULE: Often, a single text block contains multiple paragraphs separated by ` \\n `. If you see multiple images stacked vertically on one side (different Y values), and a single multi-paragraph text block on the other side, you must split the text block by `\\n` and pair each individual paragraph with its corresponding image (based on similar Y heights) inside separate {{% columns %}} blocks.
4. Rich Text Preservation: If you see markdown bold (`**text**`) or bullets (`- text`) in the raw text elements, preserve them exactly as they are in the final output.
5. Images: Render images using standard markdown syntax based on the input data. Do not alter the filenames.
6. Title Deduplication: Once a text block is extracted as the slide's `title` attribute, DO NOT repeat it inside the slide body (e.g., do NOT write `### Title` or `**Title**` at the top of the content).
7. Full-Width Interactive Components: Interactive components (e.g., `{{% textInput %}}`, `{{% quiz %}}`, etc.) and any immediate action prompt or challenge text directly preceding them (e.g., "**Your challenge:** ...") must NEVER be placed inside a `{{% column %}}`. If the slide uses `{{% columns %}}` for contextual content or images, close the `{{% columns %}}` block BEFORE rendering the action prompt and interactive component, ensuring they sit at the root slide level across the full width.

Input Data:
Process the following slides and generate the clean Markdoc code. Do not include explanations, just output the .mdoc code.

<input>
{chr(10).join(spatial_data_output)}
</input>
"""

        print("\n" + "="*60, file=sys.stderr)
        print("✅ Parsing complete! Generated prompt for LLM.", file=sys.stderr)
        print("Copy the text below (stdout) and paste it into your LLM.", file=sys.stderr)
        if schema_found:
            print("The schema was successfully retrieved from " + schema_url, file=sys.stderr)
        else:
            print("⚠️ The schema could not be retrieved; You need to manually paste it into the prompt BEFORE submitting it to the LLM.", file=sys.stderr)
        print("="*60 + "\n", file=sys.stderr)
        
        # Only print the actual prompt to stdout
        print(prompt_text)
        
        # Flush stdout to ensure the prompt prints before reading input
        sys.stdout.flush()

        # --- WAIT FOR USER INPUT ---
        print("\n" + "="*60, file=sys.stderr)
        print("📥 Paste the LLM's Markdoc response here.", file=sys.stderr)
        print("Press Enter, then either Ctrl+D (Linux/macOS) or Ctrl+Z then Enter again (Windows) when finished.", file=sys.stderr)
        print("="*60, file=sys.stderr)
        
        # Read input, remove leftover Windows EOF characters (Ctrl+Z), and strip whitespace
        llm_output = sys.stdin.read().replace('\x1a', '').strip()
        
        if not llm_output:
            print("⚠️ No output provided. Skipping post-processing for this file.", file=sys.stderr)
            continue
            
        print("\n🔄 Starting post-processing...", file=sys.stderr)
        
        # --- POST-PROCESSING ---
        
        # 1. Resolve Image Paths via Regex
        # Looks for ![filename.ext](anything_or_nothing) and replaces it with the full Astro asset path
        astro_asset_path = f"@admin-assets/images/courses/{unit_folder}"
        processed_mdoc = re.sub(
            r'!\[([^\]]+)\]\([^\)]*\)', 
            rf'![\1]({astro_asset_path}/\1)', 
            llm_output
        )

        # 2. Remove Redundant Slide Titles
        # The LLM may have repeated the slide title at the top of the slide content.
        def clean_slide(match):
            full_slide = match.group(0)
            title = match.group(1).strip()
            
            # Regex to match either titles (e.g., ### Title) or emphasized/bolded titles (e.g., **Title**) at the start of the slide content, allowing for optional whitespace and newlines
            # Regex explanation:
            # - (%}}\s*) : Captures the closing slide tag and any whitespace after it
            # - (?:#|\*)* : Non-capturing group to match any number of '#' or '* characters (for titles, emphasized or bolded text)
            # - \s*{re.escape(title)}\s* : Matches the slide title with optional whitespace around it
            # - (?:\*)* : Optionally matches closing '*' characters (any number of them, for both '*' and '**') in case the title was bolded or emphasized
            # - \n+ : Matches one or more newlines after the title
            pattern = rf'(%}}\s*)(?:#|\*)*\s*{re.escape(title)}\s*(?:\*)*\s*\n+'
            cleaned_content = re.sub(pattern, r'\1', full_slide, flags=re.IGNORECASE)
            return cleaned_content

        # Capture the entire slide block and its title for processing
        # Regex explanation:
        # - {%\s*slide\s+title="([^"]+)"[^%]*%} : Matches the opening slide tag and captures the title
        # - .*? : Non-greedy match for any content inside the slide
        # - {%\s*/slide\s*%} : Matches the closing slide tag
        slide_pattern = r'{%\s*slide\s+title="([^"]+)"[^%]*%}.*?{%\s*/slide\s*%}'
        processed_mdoc = re.sub(slide_pattern, clean_slide, processed_mdoc, flags=re.DOTALL)
        
        # 3. Append Frontmatter
        frontmatter = f"---\nunit: '{unit_str}'\ntitle: '{title_str}'\n---\n\n"
        final_mdoc = frontmatter + processed_mdoc
        
        # 3. Save the final .mdoc file
        os.makedirs(CONTENT_DIR, exist_ok=True)
        mdoc_file_path = os.path.join(CONTENT_DIR, f"{unit_folder}.mdoc")
        with open(mdoc_file_path, "w", encoding="utf-8") as f:
            f.write(final_mdoc)
        print(f"✔️ Saved final Markdoc to: {mdoc_file_path}", file=sys.stderr)
        
        # 4. Copy images to the Astro assets directory
        final_img_dir = os.path.join(ASSETS_DIR, unit_folder)
        os.makedirs(final_img_dir, exist_ok=True)
        
        copied_count = 0
        for img_file in os.listdir(output_img_dir):
            src_path = os.path.join(output_img_dir, img_file)
            dest_path = os.path.join(final_img_dir, img_file)
            shutil.copy2(src_path, dest_path)
            copied_count += 1
            
        print(f"✔️ Copied {copied_count} images to: {final_img_dir}", file=sys.stderr)
        print(f"🎉 Unit {unit_str} fully processed and integrated!", file=sys.stderr)

if __name__ == "__main__":
    extract_pptx_data()